"""Read the real production workbook with independently installed openpyxl.
Optional diagnostic; installs nothing and uses only synthetic rows.
"""
from pathlib import Path
import shutil,subprocess
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[2]
output=root/'build/task14-export/independent-workbook.xlsx'
output.parent.mkdir(parents=True,exist_ok=True)
subprocess.run([shutil.which('dart'),'run','test/platform/write_workbook_fixture.dart',str(output)],cwd=root,check=True,timeout=90)
workbook=load_workbook(output,read_only=True,data_only=False)
try:
 assert workbook.sheetnames==['Eventos']
 sheet=workbook['Eventos']
 rows=list(sheet.values)
 assert len(rows)==1206,len(rows)
 assert rows[0]==('Convertido','Congregação','Evento','Data do Evento','Registrado por')
 assert rows[1]==('João 1','São José','=1+2','01/10/2026','ação@example.test')
 assert rows[-1][0]=='João 1205'
 assert all(c.data_type=='s' for row in sheet.iter_rows() for c in row)
 print('PASS: independent openpyxl opens actual XLSX with1205 rows, exact headers/accents and literal string cells, no formulas')
finally:
 workbook.close()
